import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

def generar_xlsx(data: dict, output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resultados QA"

    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    pass_fill = PatternFill("solid", fgColor="E2EFDA")
    fail_fill = PatternFill("solid", fgColor="FCE4D6")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    # Encabezado del reporte
    ws.merge_cells("A1:F1")
    ws["A1"] = "REPORTE DE EJECUCIÓN QA"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws["A1"].alignment = center

    ws["A2"] = f"Execution ID: {data['execution_id']}"
    ws["A3"] = f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws["A4"] = f"Total: {data['total']}  |  Pasados: {data['pasados']}  |  Fallidos: {data['fallidos']}"
    ws["A5"] = ""

    # Headers tabla
    headers = ["#", "Caso", "Acción", "Esperado", "Obtenido", "Estado", "Tiempo (ms)"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin

    # Datos
    todos_los_pasos = []
    for resultado in data["resultados"]:
        for step in resultado["pasos"]:
            step["nombre_caso"] = resultado["nombre"]
            todos_los_pasos.append(step)

    for i, step in enumerate(todos_los_pasos, 1):
        row = 6 + i
        estado = step.get("estado", "").upper()
        fill = pass_fill if estado == "PASS" else fail_fill

        values = [
            i,
            step.get("nombre_caso", ""),
            step.get("entrada_accion", ""),
            step.get("resultado_esperado", ""),
            step.get("resultado_obtenido", ""),
            estado,
            step.get("tiempo_ms") or 0
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill
            cell.border = thin
            cell.alignment = center if col in [1, 6, 7] else left

    # Ancho columnas
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 40
    ws.column_dimensions["D"].width = 35
    ws.column_dimensions["E"].width = 35
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 12

    wb.save(output_path)
    return output_path