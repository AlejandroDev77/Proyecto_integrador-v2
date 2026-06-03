from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
import os
from datetime import datetime

def generar_docx(data: dict, output_path: str):
    doc = Document()

    # Título
    titulo = doc.add_heading("Reporte de Ejecución QA", 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Metadata
    doc.add_paragraph(f"Execution ID: {data['execution_id']}")
    doc.add_paragraph(f"Fecha: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    doc.add_paragraph(f"Estado: {data['estado'].upper()}")
    doc.add_paragraph(f"Total: {data['total']}  |  Pasados: {data['pasados']}  |  Fallidos: {data['fallidos']}")
    doc.add_paragraph("")

    # Tabla de resultados
    doc.add_heading("Resultados por caso", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"

    headers = ["Caso", "Resultado Esperado", "Resultado Obtenido", "Estado", "Tiempo (ms)"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True

    for step in data["steps"]:
        row = table.add_row().cells
        row[0].text = step["nombre"]
        row[1].text = step["resultado_esperado"] or ""
        row[2].text = step["resultado_obtenido"] or ""
        row[3].text = step["estado"].upper()
        row[4].text = str(step["tiempo_ms"]) if step["tiempo_ms"] else "0"

    doc.add_paragraph("")

    # Screenshots
    doc.add_heading("Evidencias (Screenshots)", level=1)
    for step in data["steps"]:
        doc.add_paragraph(f"Caso: {step['nombre']}")
        path = step.get("screenshot_path")
        if path and os.path.exists(path):
            doc.add_picture(path, width=Inches(5))
        else:
            doc.add_paragraph("(Sin screenshot disponible)")
        doc.add_paragraph("")

    doc.save(output_path)
    return output_path