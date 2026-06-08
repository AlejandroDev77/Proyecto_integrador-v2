from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
from datetime import datetime

def set_cell_bg(cell, color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def cell_text(cell, text: str, bold=False, size=10, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    para = cell.paragraphs[0]
    para.alignment = align
    run = para.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def generar_tqa711(data: dict, output_path: str):
    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    fecha_actual = datetime.now().strftime("%d/%m/%Y %H:%M:%S")

    for i, resultado in enumerate(data["resultados"]):
        if i > 0:
            doc.add_page_break()

        # ══════════════════════════════════════════
        # SECCIÓN 1 — ENCABEZADO DEL CASO (TQA-711)
        # ══════════════════════════════════════════
        titulo_para = doc.add_paragraph()
        titulo_run = titulo_para.add_run("CASO DE PRUEBA")
        titulo_run.bold = True
        titulo_run.font.size = Pt(13)
        titulo_run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        titulo_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        tabla1 = doc.add_table(rows=8, cols=2)
        tabla1.style = "Table Grid"
        tabla1.alignment = WD_TABLE_ALIGNMENT.CENTER

        # Fila 1: Título | ID
        set_cell_bg(tabla1.rows[0].cells[0], "1F4E79")
        cell_text(tabla1.rows[0].cells[0], f"Título: {resultado['nombre']}", bold=True, color="FFFFFF")
        set_cell_bg(tabla1.rows[0].cells[1], "1F4E79")
        cell_text(tabla1.rows[0].cells[1], f"ID: {resultado['case_id'][:8].upper() if resultado['case_id'] else 'AUTO'}", bold=True, color="FFFFFF")

        # Fila 2: Autor | Fecha
        set_cell_bg(tabla1.rows[1].cells[0], "BDD7EE")
        cell_text(tabla1.rows[1].cells[0], f"Autor: {resultado['autor']}", bold=True)
        set_cell_bg(tabla1.rows[1].cells[1], "BDD7EE")
        cell_text(tabla1.rows[1].cells[1], f"Fecha: {fecha_actual}", bold=True)

        # Fila 3: Caso de Uso
        tabla1.rows[2].cells[0].merge(tabla1.rows[2].cells[1])
        set_cell_bg(tabla1.rows[2].cells[0], "DEEAF1")
        cell_text(tabla1.rows[2].cells[0], f"Caso de Uso perteneciente: {resultado['caso_uso']}", bold=True)

        # Fila 4: Módulo | Versión
        set_cell_bg(tabla1.rows[3].cells[0], "BDD7EE")
        cell_text(tabla1.rows[3].cells[0], f"Módulo: {resultado['modulo']}", bold=True)
        set_cell_bg(tabla1.rows[3].cells[1], "BDD7EE")
        cell_text(tabla1.rows[3].cells[1], f"Versión del Sistema: {resultado['version_sistema']}", bold=True)

        # Fila 5: Descripción
        tabla1.rows[4].cells[0].merge(tabla1.rows[4].cells[1])
        set_cell_bg(tabla1.rows[4].cells[0], "DEEAF1")
        cell_text(tabla1.rows[4].cells[0], f"Descripción: Prueba automatizada del módulo {resultado['modulo']} — {resultado['nombre']}", bold=True)

        # Fila 6: Datos Requeridos | Prerrequisitos
        set_cell_bg(tabla1.rows[5].cells[0], "BDD7EE")
        cell_text(tabla1.rows[5].cells[0], f"Datos Requeridos: {resultado['datos_requeridos']}", bold=True)
        set_cell_bg(tabla1.rows[5].cells[1], "BDD7EE")
        cell_text(tabla1.rows[5].cells[1], f"Prerrequisitos: {resultado['prerequisitos']}", bold=True)

        # Fila 7: Postcondiciones
        tabla1.rows[6].cells[0].merge(tabla1.rows[6].cells[1])
        set_cell_bg(tabla1.rows[6].cells[0], "DEEAF1")
        cell_text(tabla1.rows[6].cells[0], f"Postcondiciones: {resultado['postcondiciones']}", bold=True)

        # Fila 8: Notas
        tabla1.rows[7].cells[0].merge(tabla1.rows[7].cells[1])
        cell_text(tabla1.rows[7].cells[0], "Notas: Caso generado automáticamente por el módulo QA del sistema.")

        doc.add_paragraph("")

        # ══════════════════════════════════════════
        # SECCIÓN 2 — EJECUCIÓN
        # ══════════════════════════════════════════
        sec2_para = doc.add_paragraph()
        sec2_run = sec2_para.add_run("EJECUCIÓN DE PRUEBA")
        sec2_run.bold = True
        sec2_run.font.size = Pt(11)
        sec2_run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

        tabla2 = doc.add_table(rows=1, cols=2)
        tabla2.style = "Table Grid"
        set_cell_bg(tabla2.rows[0].cells[0], "2E75B6")
        cell_text(tabla2.rows[0].cells[0], f"ID: {resultado['case_id'][:8].upper() if resultado['case_id'] else 'AUTO'}", bold=True, color="FFFFFF")
        set_cell_bg(tabla2.rows[0].cells[1], "2E75B6")
        cell_text(tabla2.rows[0].cells[1], f"Fecha: {fecha_actual}", bold=True, color="FFFFFF")

        resp_row = tabla2.add_row()
        resp_row.cells[0].merge(resp_row.cells[1])
        set_cell_bg(resp_row.cells[0], "DEEAF1")
        cell_text(resp_row.cells[0], "Responsable: QA Automatizado — Playwright + Python", bold=True)

        # Headers tabla ejecución
        header_row = tabla2.add_row()
        headers = ["PASO", "ENTRADA O ACCIÓN", "RESULTADO ESPERADO", "RESULTADO OBTENIDO", "DEFECTOS", "PASA/FALLA"]
        cols_tabla2 = tabla2.add_row()  # fila flujo
        cols_tabla2.cells[0].merge(cols_tabla2.cells[1])
        set_cell_bg(cols_tabla2.cells[0], "1F4E79")
        cell_text(cols_tabla2.cells[0], "FLUJO A PROBAR", bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

        tabla_exec = doc.add_table(rows=1, cols=6)
        tabla_exec.style = "Table Grid"
        for j, h in enumerate(headers):
            set_cell_bg(tabla_exec.rows[0].cells[j], "2E75B6")
            cell_text(tabla_exec.rows[0].cells[j], h, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

        for step in resultado["pasos"]:
            fila = tabla_exec.add_row()
            estado = step.get("estado", "PASA")
            bg = "E2EFDA" if estado == "PASA" else "FCE4D6"

            cell_text(fila.cells[0], str(step["paso"]), align=WD_ALIGN_PARAGRAPH.CENTER)
            cell_text(fila.cells[1], step.get("entrada_accion", ""))
            cell_text(fila.cells[2], step.get("resultado_esperado", ""))
            cell_text(fila.cells[3], step.get("resultado_obtenido", ""))
            cell_text(fila.cells[4], step.get("defectos", "—"))
            set_cell_bg(fila.cells[5], bg)
            cell_text(fila.cells[5], estado, bold=True,
                      color="375623" if estado == "PASA" else "833C00",
                      align=WD_ALIGN_PARAGRAPH.CENTER)

        doc.add_paragraph("")

        # ══════════════════════════════════════════
        # SECCIÓN 3 — EVIDENCIAS
        # ══════════════════════════════════════════
        sec3_para = doc.add_paragraph()
        sec3_run = sec3_para.add_run("EVIDENCIAS")
        sec3_run.bold = True
        sec3_run.font.size = Pt(11)
        sec3_run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)

        tabla_ev = doc.add_table(rows=1, cols=4)
        tabla_ev.style = "Table Grid"

        ev_headers = ["PASO", "ENTRADA O ACCIÓN", "OBSERVACIONES", "EVIDENCIA"]
        for j, h in enumerate(ev_headers):
            set_cell_bg(tabla_ev.rows[0].cells[j], "2E75B6")
            cell_text(tabla_ev.rows[0].cells[j], h, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

        flujo_ev = tabla_ev.add_row()
        flujo_ev.cells[0].merge(flujo_ev.cells[3])
        set_cell_bg(flujo_ev.cells[0], "1F4E79")
        cell_text(flujo_ev.cells[0], "FLUJO A PROBAR", bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)

        for step in resultado["pasos"]:
            fila = tabla_ev.add_row()
            cell_text(fila.cells[0], str(step["paso"]), align=WD_ALIGN_PARAGRAPH.CENTER)
            cell_text(fila.cells[1], step.get("entrada_accion", ""))
            cell_text(fila.cells[2], step.get("observaciones", "Ejecutado correctamente"))

            # Screenshot embebido
            screenshot = step.get("screenshot_path")
            if screenshot and os.path.exists(screenshot):
                fila.cells[3].text = ""
                para = fila.cells[3].paragraphs[0]
                run = para.add_run()
                run.add_picture(screenshot, width=Inches(2.5))
            else:
                cell_text(fila.cells[3], "(Sin evidencia)")

        doc.add_paragraph("")

        # Tiempo total
        tiempo_para = doc.add_paragraph()
        tiempo_run = tiempo_para.add_run(
            f"⏱ Tiempo total de ejecución: {resultado['tiempo_total_ms']}ms  |  "
            f"Estado: {resultado['estado_general']}  |  "
            f"Generado: {fecha_actual}"
        )
        tiempo_run.font.size = Pt(9)
        tiempo_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    doc.save(output_path)
    return output_path